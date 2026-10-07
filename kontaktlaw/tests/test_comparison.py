"""Synthetic-only regression tests. No client files or external services."""
import contextlib
import hashlib
import importlib.util
import io
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / 'skills/legal-review/scripts/compare_documents.py'
spec = importlib.util.spec_from_file_location('comparison', SCRIPT)
c = importlib.util.module_from_spec(spec)
spec.loader.exec_module(c)


def make_docx(path, paragraphs):
    from xml.sax.saxutils import escape
    body = ''.join(f'<w:p><w:r><w:t xml:space="preserve">{escape(p)}</w:t></w:r></w:p>' for p in paragraphs)
    return raw_docx(path, body)


def raw_docx(path, body, extras=None):
    with ZipFile(path, 'w') as archive:
        archive.writestr('word/document.xml', f'<w:document xmlns:w="{c.W[1:-1]}" xmlns:r="{c.R[1:-1]}"><w:body>{body}</w:body></w:document>')
        for name, data in (extras or {}).items():
            archive.writestr(name, data)
    return path


def document(*texts, complete=True, method='native'):
    result = {'spans': [], 'coverage': {'extraction_complete': complete}}
    for index, text in enumerate(texts):
        c.add_span(result, text, f'paragraph {index+1}', method)
    return result


class DifferenceTests(unittest.TestCase):
    def test_identical_and_wrapping(self):
        result = c.compare(document('1. Payment is due after acceptance.'),
                           document('1. Payment is due', 'after\nacceptance.'))
        self.assertEqual(result['changes'], [])
        self.assertEqual(result['coverage']['status'], 'no_text_changes')

    def test_addition_deletion_and_exact_offsets(self):
        for before, after, kind in [('Buyer pays.', 'Buyer pays promptly.', 'addition'),
                                    ('Buyer pays promptly.', 'Buyer pays.', 'deletion'),
                                    ('Payment is 100 AZN.', 'Payment is 200 AZN.', 'replacement')]:
            result = c.compare(document(before), document(after))
            self.assertEqual(result['changes'][0]['type'], kind)
            self.check_quotes(result)

    def check_quotes(self, result):
        for change in result['changes']:
            for side in ('before', 'after'):
                spans = {s['id']: s for s in result[side]['spans']}
                for passage in change[side]:
                    source = spans[passage['span_id']]['text']
                    self.assertEqual(passage['quote'], source[passage['start']:passage['end']])

    def test_negations_and_azerbaijani_unicode(self):
        result = c.compare(document('Alıcı ödəməlidir. Buyer shall pay 100 AZN.'),
                           document('Alıcı ödəməməlidir. Buyer shall not pay 100 AZN.'))
        self.assertGreaterEqual(len(result['changes']), 2)
        self.assertTrue(any('not' in p['quote'] for item in result['changes'] for p in item['after']))
        self.check_quotes(result)
        added = next(item for item in result['changes'] if any('not' in p['quote'] for p in item['after']))
        self.assertIn('Buyer shall pay', added['before_anchor']['context'])

    def test_unique_moved_paragraph(self):
        a = 'Buyer must deliver all goods immediately.'
        b = 'Seller shall retain all original payment records.'
        result = c.compare(document(a, b), document(b, a))
        self.assertEqual([item['type'] for item in result['changes']], ['move'])
        self.check_quotes(result)

    def test_repeated_clause_is_not_speculative_move(self):
        a = 'Buyer must deliver all goods immediately.'
        b = 'Seller shall retain all original payment records.'
        result = c.compare(document(a, b, a), document(a, a, b))
        for change in result['changes']:
            if change['type'] == 'move':
                self.assertIn('Seller', change['before'][0]['context'])

    def test_incomplete_never_means_no_changes(self):
        result = c.compare(document('Same text', complete=False), document('Same text'))
        self.assertEqual(result['coverage']['status'], 'inconclusive')
        self.assertFalse(result['coverage']['complete'])
        result = c.compare(document(), document())
        # Extractors, rather than the low-level diff, label empty input coverage.
        self.assertEqual(result['changes'], [])

    def test_ocr_change_is_uncertain(self):
        result = c.compare(document('100 AZN', method='ocr', complete=False), document('200 AZN'))
        self.assertEqual(result['changes'][0]['uncertainty'], 'ocr_requires_visual_verification')
        self.assertIsNone(result['changes'][0]['legal_effect'])

    def test_table_redistribution_is_not_false_equivalence(self):
        before = document('A', 'B C')
        after = document('A B', 'C')
        before['tables'] = [{'rows': [['A', 'B C']]}]
        after['tables'] = [{'rows': [['A B', 'C']]}]
        result = c.compare(before, after)
        self.assertEqual(result['coverage']['status'], 'inconclusive')
        self.assertTrue(result['coverage']['structural_warnings'])


class TempFiles(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)


class FileTests(TempFiles):

    def test_docx_source_unchanged_and_comments_separate(self):
        path = raw_docx(self.folder/'a.docx', '<w:p><w:r><w:t>Operative text</w:t></w:r></w:p>',
                        {'word/comments.xml': f'<w:comments xmlns:w="{c.W[1:-1]}"><w:comment w:id="1"><w:p><w:r><w:t>Negotiation only</w:t></w:r></w:p></w:comment></w:comments>'})
        before = path.read_bytes()
        result = c.extract(path)
        self.assertEqual(result['spans'][0]['text'], 'Operative text')
        self.assertEqual(result['comments'][0]['text'], 'Negotiation only')
        self.assertEqual(path.read_bytes(), before)
        self.assertEqual(result['sha256'], hashlib.sha256(before).hexdigest())

    def test_table_locations_and_header(self):
        path = raw_docx(self.folder/'a.docx', '<w:tbl><w:tr><w:tc><w:p><w:r><w:t>Amount</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>100</w:t></w:r></w:p></w:tc></w:tr></w:tbl>',
                        {'word/header1.xml': f'<w:hdr xmlns:w="{c.W[1:-1]}"><w:p><w:r><w:t>Header</w:t></w:r></w:p></w:hdr>'})
        result = c.extract(path)
        self.assertIn('table 1, row 1, cell 2', result['spans'][1]['locator'])
        self.assertIn('header1.xml', result['spans'][2]['locator'])

    def test_revision_selection(self):
        path = raw_docx(self.folder/'a.docx', '<w:p><w:r><w:t>Pay </w:t></w:r><w:del><w:r><w:delText>100</w:delText></w:r></w:del><w:ins><w:r><w:t>200</w:t></w:r></w:ins></w:p>')
        with self.assertRaisesRegex(c.ComparisonError, 'unresolved tracked changes'):
            c.extract(path)
        self.assertEqual(c.extract(path, revision='original')['spans'][0]['text'], 'Pay 100')
        self.assertEqual(c.extract(path, revision='revised')['spans'][0]['text'], 'Pay 200')

    def test_complex_revision_rejected(self):
        path = raw_docx(self.folder/'a.docx', '<w:p><w:pPr><w:pPrChange/></w:pPr><w:r><w:t>Text</w:t></w:r></w:p>')
        with self.assertRaisesRegex(c.ComparisonError, 'complex tracked changes'):
            c.extract(path, revision='revised')

    def test_cell_revision_rejected(self):
        path = raw_docx(self.folder/'a.docx', '<w:tbl><w:tr><w:tc><w:tcPr><w:cellDel/></w:tcPr><w:p><w:r><w:t>Text</w:t></w:r></w:p></w:tc></w:tr></w:tbl>')
        with self.assertRaises(c.ComparisonError):
            c.extract(path)
        with self.assertRaisesRegex(c.ComparisonError, 'complex tracked changes'):
            c.extract(path, revision='revised')

    def test_embedded_object_coverage(self):
        path = raw_docx(self.folder/'a.docx', '<w:p><w:r><w:t>Text</w:t><w:object/></w:r></w:p>')
        self.assertFalse(c.extract(path)['coverage']['extraction_complete'])

    def test_cli_new_output_and_overwrite_refusal(self):
        left = make_docx(self.folder/'a.docx', ['Pay 100 AZN.'])
        right = make_docx(self.folder/'b.docx', ['Pay 200 AZN.'])
        out = self.folder/'comparison.json'
        with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
            self.assertEqual(c.main([str(left), str(right), '--out', str(out)]), 0)
            first = out.read_bytes()
            self.assertEqual(c.main([str(left), str(right), '--out', str(out)]), 2)
            self.assertEqual(c.main([str(left), str(right), '--out', str(left)]), 2)
        self.assertEqual(out.read_bytes(), first)
        self.assertEqual(json.loads(first)['schema_version'], 'kontaktlaw-comparison/1.0')

    def test_unicode_output_path_on_legacy_windows_console(self):
        left = make_docx(self.folder/'a.docx', ['Pay 100 AZN.'])
        out = self.folder/'dəyişiklik.json'
        process = subprocess.run([sys.executable, str(SCRIPT), str(left), str(left), '--out', str(out)],
                                 capture_output=True, encoding='utf-8',
                                 env={**os.environ, 'PYTHONIOENCODING': 'cp1252'})
        self.assertEqual(process.returncode, 0, process.stderr)
        self.assertIn('dəyişiklik.json', process.stdout)

    def test_bad_file_and_unsupported_format(self):
        path = self.folder/'a.docx'
        path.write_text('broken')
        with self.assertRaises(c.ComparisonError):
            c.extract(path)
        with self.assertRaisesRegex(c.ComparisonError, 'Supported formats'):
            c.extract(self.folder/'a.doc')

    def test_empty_document_inconclusive(self):
        path = make_docx(self.folder/'empty.docx', [])
        extracted = c.extract(path)
        self.assertEqual(c.compare(extracted, extracted)['coverage']['status'], 'inconclusive')


class DependencyTests(unittest.TestCase):
    def test_missing_tesseract(self):
        with patch.dict(os.environ, {}, clear=True), patch.object(c.shutil, 'which', return_value=None):
            with self.assertRaisesRegex(c.ComparisonError, 'Tesseract is unavailable'):
                c.OCR().check()

    def test_missing_language(self):
        with patch.object(c.OCR, 'run', return_value='List of languages:\neng\n'):
            with self.assertRaisesRegex(c.ComparisonError, 'aze, rus'):
                c.OCR(executable='tesseract').check()

    def test_timeout_is_actionable(self):
        with patch.object(c.subprocess, 'run', side_effect=subprocess.TimeoutExpired('tesseract', 120)):
            with self.assertRaisesRegex(c.ComparisonError, 'exceeded'):
                c.OCR(executable='tesseract').run([])

    def test_bad_language_argument(self):
        with self.assertRaises(c.ComparisonError):
            c.OCR(languages='--help')


@unittest.skipUnless(importlib.util.find_spec('pymupdf'), 'Optional PyMuPDF is not installed')
class PDFTests(TempFiles):
    def make_pdf(self, path, text):
        import pymupdf
        with pymupdf.open() as pdf:
            page = pdf.new_page()
            page.insert_text((72, 72), text)
            pdf.save(path)
        return path

    def test_cross_format_and_line_wrap(self):
        word = make_docx(self.folder/'a.docx', ['1. Buyer shall pay 100 AZN after delivery.'])
        pdf = self.make_pdf(self.folder/'b.pdf', '1. Buyer shall pay 100 AZN\nafter delivery.')
        result = c.compare(c.extract(word), c.extract(pdf))
        self.assertEqual(result['changes'], [])
        self.assertTrue(result['coverage']['complete'])

    def test_encrypted_pdf(self):
        import pymupdf
        path = self.folder/'locked.pdf'
        with pymupdf.open() as pdf:
            pdf.new_page().insert_text((72, 72), 'Private')
            pdf.save(path, encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw='test', owner_pw='owner')
        with self.assertRaisesRegex(c.ComparisonError, 'Encrypted PDF'):
            c.extract(path)

    def test_blank_page_without_ocr(self):
        path = self.make_pdf(self.folder/'blank.pdf', '')
        result = c.extract(path, c.OCR(mode='never'))
        self.assertFalse(result['coverage']['extraction_complete'])

    def test_corrupt_pdf_cli_has_no_traceback(self):
        path = self.folder/'broken.pdf'
        path.write_bytes(b'not a pdf')
        output = io.StringIO()
        with contextlib.redirect_stderr(output):
            status = c.main([str(path), str(path), '--out', str(self.folder/'result.json')])
        self.assertEqual(status, 2)
        self.assertNotIn('Traceback', output.getvalue())


@unittest.skipUnless(os.environ.get('KONTAKTLAW_RUN_OCR_TESTS') == '1', 'Opt-in real local Tesseract tests')
class RealOCRTests(TempFiles):
    TEXTS = {'aze': 'Alıcı ödənişi 30 gün ərzində həyata keçirməlidir.',
             'eng': 'The buyer shall pay 100 AZN within 30 days.',
             'rus': 'Покупатель обязан оплатить товар в течение 30 дней.'}

    def bitmap(self, text):
        from PIL import Image, ImageDraw, ImageFont
        image = Image.new('RGB', (2000, 320), 'white')
        font_path = os.environ.get('KONTAKTLAW_TEST_FONT', 'C:/Windows/Fonts/arial.ttf')
        font = ImageFont.truetype(font_path, 48)
        draw = ImageDraw.Draw(image)
        draw.text((55, 70), text, fill='black', font=font)
        return image

    def test_three_languages_real_ocr(self):
        for language, text in self.TEXTS.items():
            with self.subTest(language=language):
                path = self.folder/f'{language}.png'
                self.bitmap(text).save(path)
                result = c.extract(path, c.OCR(languages=language, review_dir=self.folder/f'{language}-review'))
                actual = ' '.join(s['text'] for s in result['spans'])
                expected = c.TOKEN.findall(text)
                self.assertGreater(c.SequenceMatcher(None, expected, c.TOKEN.findall(actual)).ratio(), 0.9, actual)
                self.assertFalse(result['coverage']['extraction_complete'])
                self.assertTrue(Path(result['spans'][0]['ocr']['review_image']).exists())
                for span in result['spans']:
                    for word in span['ocr_words']:
                        self.assertEqual(word['text'], span['text'][word['start']:word['end']])

    def test_rotated_image(self):
        path = self.folder/'rotated.png'
        self.bitmap(self.TEXTS['eng']).rotate(90, expand=True).save(path)
        result = c.extract(path, c.OCR(languages='eng'))
        self.assertIn('100', ' '.join(s['text'] for s in result['spans']))
        self.assertFalse(result['coverage']['extraction_complete'])

    def test_default_multilingual_comparison_cli(self):
        png = self.folder/'scan.png'
        self.bitmap(self.TEXTS['aze']).save(png)
        word = make_docx(self.folder/'before.docx', [self.TEXTS['aze'].replace('30', '60')])
        out = self.folder/'comparison.json'
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(c.main([str(word), str(png), '--out', str(out)]), 0)
        result = json.loads(out.read_text(encoding='utf-8'))
        self.assertTrue(result['changes'])
        self.assertTrue(any('30' in p['quote'] for item in result['changes'] for p in item['after']))
        self.assertFalse(result['coverage']['complete'])
        self.assertTrue(out.with_name('comparison.review').is_dir())

    def test_jpeg_and_poor_quality_remain_unverified(self):
        from PIL import ImageFilter
        path = self.folder/'poor.jpg'
        self.bitmap(self.TEXTS['eng']).resize((700, 112)).filter(ImageFilter.GaussianBlur(0.7)).save(path, quality=40)
        result = c.extract(path, c.OCR(languages='eng'))
        self.assertFalse(result['coverage']['extraction_complete'])
        self.assertTrue(any(w['code'] == 'ocr_unverified' for w in result['coverage']['warnings']))

    def test_unreadable_image(self):
        from PIL import Image
        path = self.folder/'blank.png'
        Image.new('RGB', (800, 300), 'white').save(path)
        result = c.extract(path, c.OCR(languages='eng'))
        self.assertEqual(c.compare(result, result)['coverage']['status'], 'inconclusive')

    def test_scanned_pdf_and_mixed_native_image_pdf(self):
        import pymupdf
        png = self.folder/'text.png'
        self.bitmap(self.TEXTS['eng']).save(png)
        for mixed in (False, True):
            path = self.folder/f'scan-{mixed}.pdf'
            with pymupdf.open() as pdf:
                page = pdf.new_page(width=700, height=500)
                if mixed:
                    page.insert_text((40, 40), 'Native heading')
                page.insert_image(pymupdf.Rect(30, 100, 670, 210), filename=str(png))
                pdf.save(path)
            result = c.extract(path, c.OCR(languages='eng'))
            all_text = ' '.join(s['text'] for s in result['spans'])
            self.assertIn('100', all_text)
            if mixed:
                self.assertIn('Native heading', all_text)
            self.assertFalse(result['coverage']['extraction_complete'])

    def test_docx_embedded_image(self):
        from docx import Document
        png = self.folder/'text.png'
        self.bitmap(self.TEXTS['eng']).save(png)
        word = Document()
        word.add_paragraph('Native heading')
        word.add_picture(str(png))
        path = self.folder/'embedded.docx'
        word.save(path)
        result = c.extract(path, c.OCR(languages='eng'))
        self.assertIn('100', ' '.join(s['text'] for s in result['spans']))
        self.assertFalse(result['coverage']['extraction_complete'])

    def test_image_pdf_rotation_and_existing_text_layer(self):
        import pymupdf
        png = self.folder/'text.png'
        self.bitmap(self.TEXTS['eng']).save(png)
        for rotation in (0, 90):
            path = self.folder/f'layered-{rotation}.pdf'
            with pymupdf.open() as pdf:
                page = pdf.new_page(width=700, height=500)
                page.insert_image(pymupdf.Rect(30, 100, 670, 210), filename=str(png))
                # Native text outside the scan makes this a mixed page.
                page.insert_text((40, 40), 'Native heading')
                page.set_rotation(rotation)
                pdf.save(path)
            result = c.extract(path, c.OCR(languages='eng'))
            self.assertIn('100', ' '.join(s['text'] for s in result['spans']))


if __name__ == '__main__':
    unittest.main()
