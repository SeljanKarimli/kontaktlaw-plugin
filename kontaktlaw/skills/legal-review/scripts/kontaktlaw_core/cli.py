from __future__ import annotations
import argparse
import json
import sys
from pathlib import Path
from .exchange import load, save, require_valid
from .documents import extract, compare, apply_docx, clean_docx
from .workspace import library_add, approve, draft, stage_source
from . import knowledge

def report(data):
    require_valid(data)
    party = data.get('selected_party') or 'ümumi baxış; hər risk üzrə təsirlənən tərəf göstərilir'
    lines = [f'Müqavilə **{party}** maraqları baxımından təhlil edilmişdir.']
    limitations=data['coverage'].get('warnings',[])+data['coverage'].get('missing_annexes',[])
    if limitations:lines.append('Əhatə məhdudiyyəti: '+'; '.join(limitations))
    locators = {s['id']: s.get('locator', s['id']) for s in data['document']['spans']}
    for i,f in enumerate(data['findings'],1):
        if data['language'] == 'az' and f.get('document_language', 'az') != 'az' and not all(f.get(k) for k in ('quote_translation','proposal_translation')):
            raise ValueError('Foreign-language findings require Azerbaijani quote and proposal translations.')
        for key in ('title','explanation','legal_basis','proposal','severity'):
            if not f.get(key): raise ValueError(f'Finding {f["id"]}: missing report field {key}')
        quote = f['quote']
        proposal = f['proposal']
        if f.get('document_language', 'az') != 'az' and data['language'] == 'az':
            quote = f'Original mətn: {quote}\n\nAzərbaycan dilinə tərcümə: {f["quote_translation"]}'
            proposal = f'Original dildə təklif: {proposal}\n\nAzərbaycan dilinə tərcümə: {f["proposal_translation"]}'
        lines.extend([f'### {i}. {f["title"]}', f'**Problemli bənd:** {locators[f["span_id"]]}',
                      f'**Problemli mətn:** {quote}', f'**Riskin izahı:** {f["explanation"]}',
                      f'**Hüquqi əsas:** {f["legal_basis"]}', f'**Qısa düzəliş təklifi:** {proposal}',
                      f'**Risk səviyyəsi:** {f["severity"]}'])
        if i < len(data['findings']): lines.append('---')
    if not data['findings']: lines.append('Təsdiqlənmiş risk qeydi yoxdur. Bu, müqavilənin risksiz olduğu demək deyil.')
    return '\n\n'.join(lines)+'\n'

def main(argv=None):
    parser = argparse.ArgumentParser(description='KontaktLaw local evidence and document tools')
    commands = parser.add_subparsers(dest='command', required=True)
    p = commands.add_parser('extract'); p.add_argument('input'); p.add_argument('--out',required=True); p.add_argument('--ocr', action='store_true'); p.add_argument('--ocr-languages',default='aze+eng+rus')
    p = commands.add_parser('validate'); p.add_argument('input'); p.add_argument('--original')
    p = commands.add_parser('compare'); p.add_argument('before'); p.add_argument('after'); p.add_argument('--out',required=True)
    p = commands.add_parser('apply-docx'); p.add_argument('source'); p.add_argument('result'); p.add_argument('--out',required=True); p.add_argument('--edit-id',action='append'); p.add_argument('--all-grammar',action='store_true')
    p = commands.add_parser('clean-docx'); p.add_argument('source'); p.add_argument('--out',required=True); p.add_argument('--revision-view',choices=('original','revised'),required=True)
    p = commands.add_parser('report'); p.add_argument('input'); p.add_argument('--out',required=True)
    p = commands.add_parser('obligations'); p.add_argument('input'); p.add_argument('--out',required=True)
    p = commands.add_parser('article'); p.add_argument('--law',required=True); p.add_argument('--article',required=True)
    p = commands.add_parser('library-add'); p.add_argument('record'); p.add_argument('--workspace',required=True)
    p = commands.add_parser('library-approve'); p.add_argument('id'); p.add_argument('--workspace',required=True); p.add_argument('--approved-by',required=True)
    p = commands.add_parser('draft'); p.add_argument('id'); p.add_argument('--workspace',required=True); p.add_argument('--values',required=True); p.add_argument('--out',required=True)
    p = commands.add_parser('source-refresh'); p.add_argument('--source-file',required=True); p.add_argument('--candidate',required=True); p.add_argument('--workspace',required=True); p.add_argument('--official-url',required=True); p.add_argument('--retrieved-at',required=True)
    p = commands.add_parser('website-result'); p.add_argument('input'); p.add_argument('--original'); p.add_argument('--model',required=True); p.add_argument('--reasoning',required=True); p.add_argument('--prompt-fingerprint',required=True); p.add_argument('--out',required=True)
    args=parser.parse_args(argv)
    try:
        if args.command=='website-result':
            from .website import export_review
            result=export_review(load(args.input),args.model,args.reasoning,args.prompt_fingerprint,load(args.original) if args.original else None);save(args.out,result)
        elif args.command=='extract': result=extract(args.input,args.ocr,args.ocr_languages); save(args.out,result)
        elif args.command=='validate':
            require_valid(load(args.input),load(args.original) if args.original else None)
            result={'valid':True,'legal_accuracy':'Not established by structural validation.'}
        elif args.command=='compare': result=compare(load(args.before),load(args.after)); save(args.out,result)
        elif args.command=='apply-docx':
            data=load(args.result)
            selected=[edit['id'] for edit in data['edits'] if edit['mode']=='grammar'] if args.all_grammar else args.edit_id
            if selected is None: raise ValueError('Select --all-grammar or --edit-id')
            result=apply_docx(args.source,data,args.out,selected,'all' if args.all_grammar else 'conservative')
        elif args.command=='clean-docx': result=clean_docx(args.source,args.out,args.revision_view)
        elif args.command=='report':
            content=report(load(args.input)); target=Path(args.out)
            if target.exists(): raise ValueError('Output already exists')
            target.parent.mkdir(parents=True,exist_ok=True)
            if target.suffix.lower()=='.docx':
                from docx import Document
                doc=Document()
                for part in content.split('\n\n'):
                    if part.startswith('### '): doc.add_heading(part[4:],level=2)
                    elif part!='---': doc.add_paragraph(part.replace('**',''))
                doc.save(target)
            else: target.write_text(content,encoding='utf-8')
            result={'output':str(target),'visual_verification':'Required before delivering a DOCX.'}
        elif args.command=='obligations':
            data=require_valid(load(args.input)); result={'document_sha256':data['document']['sha256'],'obligations':data['obligations']};save(args.out,result)
        elif args.command=='article': result=knowledge.article(args.law,args.article)
        elif args.command=='library-add': result=library_add(args.workspace,load(args.record))
        elif args.command=='library-approve': result=approve(args.workspace,args.id,args.approved_by)
        elif args.command=='draft': result=draft(args.workspace,args.id,load(args.values));save(args.out,result)
        elif args.command=='source-refresh': result=stage_source(knowledge.ROOT,args.source_file,args.candidate,args.workspace,args.official_url,args.retrieved_at)
        print(json.dumps(result,ensure_ascii=False,indent=2))
    except (ValueError,OSError,KeyError,ImportError) as e:
        parser.exit(2,str(e)+'\n')

if __name__=='__main__':
    if hasattr(sys.stdout,'reconfigure'): sys.stdout.reconfigure(encoding='utf-8')
    main()
