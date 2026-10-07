"""Optional real-format integration tests. Run with requirements-document.txt installed."""
from pathlib import Path
import importlib.util
import sys,tempfile,unittest
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'shared'))
from kontaktlaw_core.documents import extract,apply_docx
from kontaktlaw_core.exchange import validate

class MultilingualTests(unittest.TestCase):
 def test_azerbaijani_and_russian_exact_text(self):
  for text in ('1. Alıcı ödənişi qəbuldan əvvəl etməli deyil.','1. Покупатель не обязан платить до приёмки.'):
   with self.subTest(text=text),tempfile.TemporaryDirectory() as d:
    p=Path(d)/'contract.txt';p.write_text(text,encoding='utf-8');x=extract(p);self.assertEqual(x['document']['spans'][0]['text'],text);self.assertEqual(x['document']['spans'][0]['clause_id'],'1');self.assertEqual(validate(x),[])

@unittest.skipUnless(importlib.util.find_spec('pymupdf'),'Install requirements-document.txt for PDF integration tests')
class PdfTests(unittest.TestCase):
 def test_text_and_blank_page_coverage(self):
  import pymupdf as fitz
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'mixed.pdf';pdf=fitz.open();page=pdf.new_page();page.insert_text((72,72),'Payment is due after acceptance.');pdf.new_page();pdf.save(p);pdf.close()
   x=extract(p);self.assertEqual(len(x['document']['spans']),2);self.assertFalse(x['coverage']['extraction_complete']);self.assertIn('Page 2',x['coverage']['warnings'][0]);self.assertEqual(x['document']['spans'][0]['page'],1)

@unittest.skipUnless(importlib.util.find_spec('docx'),'Install requirements-document.txt for DOCX integration tests')
class RealDocxTests(unittest.TestCase):
 def test_all_grammar_cross_run_and_deletion_preserve_package(self):
  from docx import Document
  with tempfile.TemporaryDirectory() as d:
   source=Path(d)/'source.docx';out=Path(d)/'out.docx';doc=Document();p=doc.add_paragraph()
   p.add_run('He ar').bold=True;p.add_run('e ready ready.').italic=True
   doc.sections[0].header.paragraphs[0].text='Synthetic header';doc.add_table(rows=1,cols=1).cell(0,0).text='100 AZN';doc.save(source)
   original=source.read_bytes();x=extract(source)
   x['edits']=[{'id':'e1','span_id':'p0','start':3,'end':6,'quote':'are','replacement':'is','mode':'grammar'},
               {'id':'e2','span_id':'p0','start':7,'end':13,'quote':'ready ','replacement':'','mode':'grammar'}]
   apply_docx(source,x,out,['e1','e2'],'all');reopened=Document(out)
   self.assertEqual(reopened.paragraphs[0].text,'He is ready.')
   self.assertTrue(reopened.paragraphs[0].runs[0].bold);self.assertTrue(reopened.paragraphs[0].runs[1].italic)
   self.assertEqual(reopened.sections[0].header.paragraphs[0].text,'Synthetic header')
   self.assertEqual(reopened.tables[0].cell(0,0).text,'100 AZN');self.assertEqual(source.read_bytes(),original)

 def test_save_and_reopen_preserves_table_and_run_format(self):
  from docx import Document
  with tempfile.TemporaryDirectory() as d:
   source=Path(d)/'source.docx';out=Path(d)/'out.docx';doc=Document();p=doc.add_paragraph();p.add_run('The Supplier are ready.').bold=True;doc.add_table(rows=1,cols=1).cell(0,0).text='AZN 100';doc.save(source)
   x=extract(source);x['edits']=[{'id':'e1','span_id':'p0','start':13,'end':16,'quote':'are','replacement':'is','mode':'grammar'}];apply_docx(source,x,out,['e1'])
   reopened=Document(out);self.assertEqual(reopened.paragraphs[0].text,'The Supplier is ready.');self.assertTrue(reopened.paragraphs[0].runs[0].bold);self.assertEqual(reopened.tables[0].cell(0,0).text,'AZN 100')

if __name__=='__main__':unittest.main()
